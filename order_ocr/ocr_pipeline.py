"""訂單 OCR 流程：AI 擷取 → 主檔與規則驗證 → 針對問題重新判讀 → 自動修正 → 輸出 ERP 匯入檔／人工確認清單。

用法：
  py ocr_pipeline.py samples                  # 零成本：Windows 內建 OCR（預設）
  py ocr_pipeline.py samples --engine claude  # 呼叫 Claude API（需設定 ANTHROPIC_API_KEY，會產生費用）
  py ocr_pipeline.py samples --engine mock    # 模擬模式：只測試流程是否能跑通
"""
import argparse
import base64
import csv
import difflib
import json
import random
import re
from datetime import date
from pathlib import Path

ROOT = Path(__file__).parent
MODEL = "claude-opus-5-5"

ORDER_SCHEMA = {
    "type": "object",
    "properties": {
        "customer_name": {"type": "string"},
        "po_number": {"type": "string"},
        "order_date": {"type": "string", "description": "西元 YYYY-MM-DD"},
        "delivery_date": {"type": "string", "description": "西元 YYYY-MM-DD"},
        "lines": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "item_code": {"type": "string"},
                    "description": {"type": "string"},
                    "quantity": {"type": "number"},
                    "unit_price": {"type": "number"},
                    "amount": {"type": "number"},
                },
                "required": ["item_code", "description", "quantity", "unit_price", "amount"],
                "additionalProperties": False,
            },
        },
        "total_amount": {"type": "number"},
        "uncertain_fields": {"type": "array", "items": {"type": "string"}},
    },
    "required": ["customer_name", "po_number", "order_date", "delivery_date", "lines", "total_amount", "uncertain_fields"],
    "additionalProperties": False,
}

SYSTEM_PROMPT = """你是製造業訂單處理部門的資料擷取助手。你會收到一張掃描或傳真的採購單影像，請把內容轉成結構化資料。
規則：
- 日期一律輸出西元 YYYY-MM-DD；民國年請加 1911 換算。
- 數字去掉千分位與幣別符號，只輸出數值。
- 逐字照抄料號，注意 0/O、1/I、5/S、8/B 等容易混淆的字元。
- 只擷取影像上確實存在的品項，不要自行補上。
- 印章、污漬或模糊導致看不清楚的欄位，仍輸出你最有把握的判讀，並把欄位名稱（例如 lines[2].quantity）列入 uncertain_fields。"""


# ---------------------------------------------------------------- 主檔
def load_master(data_dir: Path = ROOT / "data") -> dict:
    with open(data_dir / "master_customers.csv", encoding="utf-8-sig") as fp:
        customers = {r["customer_name"]: r["customer_code"] for r in csv.DictReader(fp)}
    with open(data_dir / "master_items.csv", encoding="utf-8-sig") as fp:
        items = {r["item_code"]: {"description": r["description"], "std_price": float(r["std_price"])} for r in csv.DictReader(fp)}
    return {"customers": customers, "items": items}


def master_context(master: dict) -> str:
    cust = "\n".join(f"- {n}（{c}）" for n, c in master["customers"].items())
    items = "\n".join(f"- {c}：{v['description']}" for c, v in master["items"].items())
    return f"""以下是公司主檔，供你核對判讀結果。料號與客戶名稱應該會出現在主檔中；若影像內容與主檔都對不上，請照影像抄寫並列入 uncertain_fields，不要硬套主檔。
【客戶主檔】
{cust}
【料號主檔】
{items}"""


# ---------------------------------------------------------------- AI 擷取
class ClaudeExtractor:
    def __init__(self, effort: str = "medium"):
        import anthropic
        self.client = anthropic.Anthropic()
        self.effort = effort
        self.calls = 0

    def extract(self, image_path: Path, extra_context: str = "", previous: dict | None = None, issues: list | None = None) -> dict:
        img_b64 = base64.standard_b64encode(image_path.read_bytes()).decode()
        text = "請擷取這張採購單的內容。"
        if extra_context:
            text += "\n\n" + extra_context
        if previous is not None:
            text += ("\n\n前一次的擷取結果如下：\n" + json.dumps(previous, ensure_ascii=False)
                     + "\n\n系統檢查發現以下問題，請回到影像中對應的位置重新仔細判讀，再輸出修正後的完整結果：\n"
                     + "\n".join(f"- {i['field']}：{i['message']}" for i in issues))
        response = self.client.beta.messages.create(
            model=MODEL,
            max_tokens=16000,
            betas=["server-side-fallback-2026-07-01"],
            fallbacks="default",  # 遇到安全分類器誤判拒答時，由伺服器自動改用備援模型
            output_config={"effort": self.effort, "format": {"type": "json_schema", "schema": ORDER_SCHEMA}},
            system=SYSTEM_PROMPT,
            messages=[{"role": "user", "content": [
                {"type": "image", "source": {"type": "base64", "media_type": "image/png", "data": img_b64}},
                {"type": "text", "text": text},
            ]}],
        )
        self.calls += 1
        if response.stop_reason == "refusal":
            raise RuntimeError(f"{image_path.name}：模型拒絕處理")
        if response.stop_reason == "max_tokens":
            raise RuntimeError(f"{image_path.name}：輸出被截斷")
        return json.loads(next(b.text for b in response.content if b.type == "text"))


class MockExtractor:
    """不呼叫 API 的模擬擷取器：以正確答案為底，依單據難度隨機加入常見 OCR 錯誤。
    只用來測試流程是否能跑通，產生的準確率數字不代表真實效果。"""

    SWAPS = {"0": "O", "1": "I", "5": "S", "8": "B", "O": "0", "S": "5", "B": "8"}
    RATE = {"clean": 0.02, "stamp": 0.05, "noisy": 0.10, "hard": 0.18}

    def __init__(self, truth_file: Path, seed: int = 7):
        self.truth = {t["file"]: t for t in json.loads(truth_file.read_text(encoding="utf-8"))}
        self.rng = random.Random(seed)
        self.calls = 0

    def _garble_code(self, code: str) -> str:
        idx = [i for i, ch in enumerate(code) if ch in self.SWAPS]
        if not idx:
            return code
        i = self.rng.choice(idx)
        return code[:i] + self.SWAPS[code[i]] + code[i + 1:]

    def extract(self, image_path: Path, extra_context: str = "", previous=None, issues=None) -> dict:
        self.calls += 1
        t = self.truth[image_path.name]
        p = self.RATE[t["difficulty"]]
        hit = lambda: self.rng.random() < p
        out = {
            "customer_name": t["customer_name"][:-2] if hit() else t["customer_name"],
            "po_number": self._garble_code(t["po_number"]) if hit() else t["po_number"],
            "order_date": t["order_date"].replace("2026", "2025") if hit() else t["order_date"],
            "delivery_date": t["delivery_date"],
            "lines": [],
            "total_amount": t["total_amount"] * (10 if hit() else 1),
            "uncertain_fields": [],
        }
        for l in t["lines"]:
            l = dict(l)
            if hit():
                l["item_code"] = self._garble_code(l["item_code"])
            if hit():
                l["quantity"] = int(str(int(l["quantity"]))[:-1] or 0)
            out["lines"].append(l)
        return out


# ---------------------------------------------------------------- 驗證與修正
# 採購單號格式：實務上可依客戶各自的單號規則設定
MAX_LEAD_DAYS = 90  # 一般訂單的交期上限，可依產品別調整
PO_PATTERN = re.compile(r"^PO\d{8}$")
CONFUSABLE = str.maketrans({"O": "0", "I": "1", "L": "1", "S": "5", "B": "8", "Z": "2"})


def canon(code: str) -> str:
    return re.sub(r"[\s_\-–—]", "", code.upper()).translate(CONFUSABLE)


def parse_date(s: str):
    try:
        return date.fromisoformat(s)
    except (TypeError, ValueError):
        return None


def validate(order: dict, master: dict) -> list[dict]:
    """回傳問題清單。level = error（必須處理）／warning（提醒）。"""
    issues = []
    add = lambda field, msg, level="error": issues.append({"field": field, "message": msg, "level": level})

    if order["customer_name"] not in master["customers"]:
        near = difflib.get_close_matches(order["customer_name"], master["customers"], n=1, cutoff=0.5)
        add("customer_name", f"客戶「{order['customer_name']}」不在客戶主檔" + (f"，最接近：{near[0]}" if near else ""))
    if not order["po_number"].strip():
        add("po_number", "缺少採購單號")
    elif not PO_PATTERN.match(order["po_number"]):
        add("po_number", f"採購單號「{order['po_number']}」不符合格式 PO + 8 碼數字")

    od, dd = parse_date(order["order_date"]), parse_date(order["delivery_date"])
    if od is None:
        add("order_date", f"訂購日期格式錯誤：{order['order_date']}")
    if dd is None:
        add("delivery_date", f"交期格式錯誤：{order['delivery_date']}")
    if od and dd and dd < od:
        add("delivery_date", "交期早於訂購日期")
    elif od and dd and (dd - od).days > MAX_LEAD_DAYS:
        add("order_date", f"訂購日到交期相隔 {(dd - od).days} 天，超過 {MAX_LEAD_DAYS} 天，年份可能誤讀")

    if not order["lines"]:
        add("lines", "沒有任何品項")
    for n, l in enumerate(order["lines"]):
        f = f"lines[{n}]"
        item = master["items"].get(l["item_code"])
        if item is None:
            add(f"{f}.item_code", f"料號「{l['item_code']}」不在料號主檔")
        elif l["unit_price"] and abs(l["unit_price"] - item["std_price"]) / item["std_price"] > 0.3:
            add(f"{f}.unit_price", f"單價 {l['unit_price']} 與標準價 {item['std_price']} 差異超過 30%", "warning")
        if l["quantity"] <= 0 or l["quantity"] != int(l["quantity"]):
            add(f"{f}.quantity", f"數量異常：{l['quantity']}")
        if abs(l["quantity"] * l["unit_price"] - l["amount"]) > 0.5:
            add(f"{f}.amount", f"數量 × 單價（{l['quantity']} × {l['unit_price']}）≠ 金額 {l['amount']}")

    line_sum = round(sum(l["amount"] for l in order["lines"]), 2)
    if abs(line_sum - order["total_amount"]) > 0.5:
        add("total_amount", f"明細金額加總 {line_sum:,.2f} ≠ 合計 {order['total_amount']:,.2f}")
    return issues


def best_match(text: str, choices, cutoff: float, margin: float = 0.08):
    """模糊比對：最佳候選要夠像，而且要明顯勝過第二名，才視為唯一答案。"""
    scored = sorted(((difflib.SequenceMatcher(None, text, c).ratio(), c) for c in choices), reverse=True)
    if not scored or scored[0][0] < cutoff:
        return None
    if len(scored) > 1 and scored[0][0] - scored[1][0] < margin:
        return None
    return scored[0][1]


def price_ok(price: float, std: float | None) -> bool:
    return std is not None and price > 0 and abs(price - std) / std <= 0.3


def fix_line(l: dict, master: dict, n: int, notes: list):
    items = master["items"]
    f = f"lines[{n}]"
    # 1) 料號：先做易混淆字元比對，再做模糊比對（需單價佐證）
    if l["item_code"] not in items:
        exact = [c for c in items if canon(c) == canon(l["item_code"])]
        if len(exact) == 1:
            notes.append(f"{f}.item_code：{l['item_code']} → {exact[0]}（易混淆字元校正）")
            l["item_code"] = exact[0]
        else:
            cands = {c: canon(c) for c in items}
            guess = best_match(canon(l["item_code"]), cands.values(), cutoff=0.7)
            code = next((c for c, v in cands.items() if v == guess), None)
            implied = l["amount"] / l["quantity"] if l["quantity"] else l["unit_price"]
            if code and (price_ok(l["unit_price"], items[code]["std_price"]) or price_ok(implied, items[code]["std_price"])):
                notes.append(f"{f}.item_code：{l['item_code']} → {code}（模糊比對＋單價佐證）")
                l["item_code"] = code
    if l["item_code"] in items:
        l["description"] = items[l["item_code"]]["description"]


def line_candidates(l: dict, master: dict) -> list[tuple[str, float, str]]:
    """數量 × 單價 ≠ 金額時，列出「假設某一欄讀錯」的所有可能修正：(欄位, 新值, 理由)。"""
    std = master["items"].get(l["item_code"], {}).get("std_price")
    q, p, a = l["quantity"], l["unit_price"], l["amount"]
    out = []
    if q > 0 and a > 0:
        p2 = round(a / q, 2)
        if abs(p2 * q - a) <= 0.5 and price_ok(p2, std):
            out.append(("unit_price", p2, "由金額 ÷ 數量回推，符合標準價"))
    if price_ok(p, std):
        if a > 0 and abs(a / p - round(a / p)) < 1e-6 and round(a / p) > 0:
            out.append(("quantity", round(a / p), "由金額 ÷ 單價回推"))
        if q > 0 and q == int(q):
            out.append(("amount", round(q * p, 2), "由數量 × 單價重算"))
    return out


def inconsistent(l: dict) -> bool:
    return abs(l["quantity"] * l["unit_price"] - l["amount"]) > 0.5


def auto_fix(order: dict, master: dict) -> list[str]:
    """只做「有唯一答案、且有其他欄位佐證」的修正，其餘交給人工。回傳修正紀錄。"""
    notes = []
    po = order["po_number"]
    if not PO_PATTERN.match(po) and po[:1].upper() == "P":
        fixed = "PO" + po[2:].upper().translate(CONFUSABLE)
        if PO_PATTERN.match(fixed):
            notes.append(f"po_number：{po} → {fixed}（易混淆字元校正）")
            order["po_number"] = fixed
    lines = order["lines"]
    for n, l in enumerate(lines):
        fix_line(l, master, n, notes)

    def apply(n, field, value, why):
        notes.append(f"lines[{n}].{field}：{lines[n][field]} → {value}（{why}）")
        lines[n][field] = value

    # 單價明顯不合理、且「金額 ÷ 數量」剛好符合標準價 → 單價讀錯的證據很強，直接修正
    for n, l in enumerate(lines):
        std = master["items"].get(l["item_code"], {}).get("std_price")
        if inconsistent(l) and not price_ok(l["unit_price"], std):
            c = [x for x in line_candidates(l, master) if x[0] == "unit_price"]
            if c:
                apply(n, *c[0])

    # 只剩一個品項對不上時，用單據上的「合計」判斷哪一欄讀錯：修正後加總必須等於合計
    bad = [n for n, l in enumerate(lines) if inconsistent(l)]
    if len(bad) == 1:
        n = bad[0]
        others = sum(l["amount"] for i, l in enumerate(lines) if i != n)
        match = [c for c in line_candidates(lines[n], master)
                 if abs(others + (c[1] if c[0] == "amount" else lines[n]["amount"]) - order["total_amount"]) <= 0.5]
        if len(match) == 1:
            apply(n, match[0][0], match[0][1], match[0][2] + "，並與合計交叉驗證")

    # 明細全部正確、只有合計對不上 → 合計讀錯，以明細加總為準
    line_sum = round(sum(l["amount"] for l in lines), 2)
    if abs(line_sum - order["total_amount"]) > 0.5 and not validate_lines_only(order, master):
        notes.append(f"total_amount：{order['total_amount']:,.2f} → {line_sum:,.2f}（明細皆通過檢查，以明細加總為準）")
        order["total_amount"] = line_sum
    # 客戶名稱不自動改寫：相似但不一致時，由 validate() 列出最接近的主檔名稱，送人工確認
    return notes


def validate_lines_only(order: dict, master: dict) -> list:
    return [i for i in validate(order, master) if i["field"].startswith("lines") and i["level"] == "error"]


# ---------------------------------------------------------------- 流程
def process(image: Path, extractor, master: dict, mode: str = "pipeline") -> dict:
    """mode="baseline"：單次擷取，不做任何檢查（對照組）。
    mode="pipeline"：提供主檔 → 驗證 → 有問題則請 AI 重新判讀 → 安全自動修正 → 判定放行或人工確認。"""
    if mode == "baseline":
        order = extractor.extract(image)
        return {"file": image.name, "order": order, "issues": [], "fixes": [], "status": "未檢查"}

    ctx = master_context(master)
    order = extractor.extract(image, extra_context=ctx)
    issues = [i for i in validate(order, master) if i["level"] == "error"]
    rechecked = False
    if issues and getattr(extractor, "supports_recheck", True):
        order = extractor.extract(image, extra_context=ctx, previous=order, issues=issues)
        rechecked = True
    fixes = auto_fix(order, master)
    issues = validate(order, master)
    errors = [i for i in issues if i["level"] == "error"]
    status = "待人工確認" if errors or order["uncertain_fields"] else "自動放行"
    return {"file": image.name, "order": order, "issues": issues, "fixes": fixes,
            "rechecked": rechecked, "status": status}


# ---------------------------------------------------------------- 輸出
def export(results: list[dict], master: dict, out: Path):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    head = wb.active
    head.title = "訂單抬頭"
    head.append(["來源檔案", "訂單類型", "客戶編號", "客戶名稱", "客戶PO", "訂購日期", "需求交期", "合計金額", "處理狀態"])
    lines = wb.create_sheet("訂單明細")
    lines.append(["客戶PO", "項次", "料號", "品名", "數量", "單價", "金額"])
    review = wb.create_sheet("人工確認清單")
    review.append(["來源檔案", "欄位", "問題", "等級"])
    log = wb.create_sheet("自動修正紀錄")
    log.append(["來源檔案", "修正內容"])

    for r in results:
        o = r["order"]
        head.append([r["file"], "OR", master["customers"].get(o["customer_name"], ""), o["customer_name"], o["po_number"],
                     o["order_date"], o["delivery_date"], o["total_amount"], r["status"]])
        for n, l in enumerate(o["lines"], 1):
            lines.append([o["po_number"], n * 10, l["item_code"], l["description"], l["quantity"], l["unit_price"], l["amount"]])
        for i in r["issues"]:
            review.append([r["file"], i["field"], i["message"], i["level"]])
        for f in o.get("uncertain_fields", []):
            review.append([r["file"], f, "AI 標記為不確定", "warning"])
        for f in r["fixes"]:
            log.append([r["file"], f])

    for ws in wb.worksheets:
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
        for col in ws.columns:
            ws.column_dimensions[col[0].column_letter].width = max(10, min(60, max(len(str(c.value or "")) for c in col) * 1.6))
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("folder", type=Path, help="放採購單影像的資料夾")
    ap.add_argument("--out", type=Path, default=ROOT / "output" / "erp_import.xlsx")
    ap.add_argument("--engine", default="local", choices=["local", "claude", "mock"])
    ap.add_argument("--effort", default="medium", choices=["low", "medium", "high", "xhigh", "max"])
    args = ap.parse_args()

    master = load_master()
    if args.engine == "local":
        from local_ocr import LocalExtractor
        extractor = LocalExtractor()
    elif args.engine == "claude":
        extractor = ClaudeExtractor(args.effort)
    else:
        extractor = MockExtractor(args.folder / "ground_truth.json")
    results = []
    for img in sorted(args.folder.glob("*.png")):
        r = process(img, extractor, master)
        print(f"{r['file']}: {r['status']}" + (f"（自動修正 {len(r['fixes'])} 處）" if r["fixes"] else ""))
        results.append(r)
    export(results, master, args.out)
    print(f"\n完成 {len(results)} 張，API 呼叫 {extractor.calls} 次 → {args.out}")


if __name__ == "__main__":
    main()
